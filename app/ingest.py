"""Ingestion pipeline: file/text -> chunks -> embeddings -> Qdrant."""
import io
import logging

from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from .config import settings
from .embeddings import embed_texts
from .vector_store import Chunk, VectorStore

logger = logging.getLogger(__name__)


def extract_text(filename: str, content: bytes) -> str:
    """Extract plain text from a file, dispatching on extension.

    Supported formats:
      .pdf  — PyPDF (text-layer extraction)
      .docx — python-docx (paragraphs + table cells)
      .xlsx — openpyxl (all cell values, sheet by sheet)
      .csv / .txt / .md — raw UTF-8 / latin-1 decode
    """
    name = filename.lower()

    # ── PDF ──────────────────────────────────────────────────────────
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n\n".join(pages)
        if not text.strip():
            logger.warning("[ingest] PDF '%s' produced no text — may be scanned/image-only.", filename)
        return text

    # ── Word (.docx) ─────────────────────────────────────────────────
    if name.endswith(".docx"):
        try:
            import docx  # python-docx
            doc = docx.Document(io.BytesIO(content))
            parts: list[str] = []
            for para in doc.paragraphs:
                if para.text.strip():
                    parts.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n\n".join(parts)
            if not text.strip():
                logger.warning("[ingest] DOCX '%s' produced no text — file may be empty.", filename)
            return text
        except Exception as e:
            logger.error("[ingest] Failed to parse DOCX '%s': %s", filename, e)
            return ""

    # ── Excel (.xlsx) ────────────────────────────────────────────────
    if name.endswith(".xlsx"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            parts: list[str] = []
            for sheet in wb.worksheets:
                parts.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None and str(c).strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n\n".join(parts)
            if not text.strip():
                logger.warning("[ingest] XLSX '%s' produced no text — file may be empty.", filename)
            return text
        except Exception as e:
            logger.error("[ingest] Failed to parse XLSX '%s': %s", filename, e)
            return ""

    # ── Plain text / Markdown / CSV ──────────────────────────────────
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1", errors="ignore")


def chunk_text(text: str) -> list[str]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    return [c for c in splitter.split_text(text) if c.strip()]


def _embed_chunks(chunks: list[str]) -> list[list[float]]:
    """Embed chunks as plain text, matching how queries are embedded.

    Previously this prepended '[Document: filename]' to each chunk before
    embedding. That caused a vector-space mismatch: stored vectors included
    the filename prefix but query vectors did not, so even an exact FAQ
    question scored poorly against its own answer. Embedding both sides the
    same way (plain text) gives consistent cosine similarity and dramatically
    improves recall for FAQ-style knowledge bases.
    """
    return embed_texts(chunks)


def ingest_file(
    store: VectorStore,
    tenant_id: str,
    knowledge_base_id: str,
    filename: str,
    content: bytes,
) -> dict:
    text = extract_text(filename, content)
    chunks = chunk_text(text)
    if not chunks:
        logger.error(
            "[ingest] '%s' produced 0 chunks (extracted %d chars). "
            "File may be unsupported format, empty, or image-only PDF.",
            filename, len(text),
        )
        return {"chunks": 0, "stored": 0}
    vectors = _embed_chunks(chunks)
    payload_chunks = [
        Chunk(
            text=t,
            tenant_id=tenant_id,
            knowledge_base_id=knowledge_base_id,
            source=filename,
            chunk_index=i,
        )
        for i, t in enumerate(chunks)
    ]
    stored = store.upsert(vectors, payload_chunks)
    return {"chunks": len(chunks), "stored": stored, "source": filename}


def ingest_text(
    store: VectorStore,
    tenant_id: str,
    knowledge_base_id: str,
    source: str,
    text: str,
) -> dict:
    chunks = chunk_text(text)
    if not chunks:
        return {"chunks": 0, "stored": 0}
    vectors = _embed_chunks(chunks)
    payload_chunks = [
        Chunk(
            text=t,
            tenant_id=tenant_id,
            knowledge_base_id=knowledge_base_id,
            source=source,
            chunk_index=i,
        )
        for i, t in enumerate(chunks)
    ]
    stored = store.upsert(vectors, payload_chunks)
    return {"chunks": len(chunks), "stored": stored, "source": source}
